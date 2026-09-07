import io
import os
import openpyxl
from fastapi import APIRouter, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Optional, List, Dict, Any
from backend.db.database import (
    get_glossary_terms,
    add_glossary_term,
    update_glossary_term,
    delete_glossary_term
)

router = APIRouter(prefix="/api/glossary", tags=["glossary"])

class GlossaryItemModel(BaseModel):
    source_term: str
    target_term: str
    target_lang: Optional[str] = "en"
    category: Optional[str] = "通用"
    is_enabled: Optional[int] = 1

@router.get("")
def list_glossary(target_lang: Optional[str] = None, category: Optional[str] = None, keyword: Optional[str] = None):
    terms = get_glossary_terms(target_lang=target_lang, category=category, keyword=keyword)
    return {
        "status": "success",
        "total": len(terms),
        "data": terms
    }

@router.post("")
def create_term(data: GlossaryItemModel):
    if not data.source_term or not data.target_term:
        raise HTTPException(status_code=400, detail="中文术语与外文译文均不能为空")
    new_id = add_glossary_term(
        source_term=data.source_term,
        target_term=data.target_term,
        target_lang=data.target_lang or "en",
        category=data.category or "通用"
    )
    return {"status": "success", "id": new_id, "message": "词条已添加"}

@router.put("/{term_id}")
def edit_term(term_id: int, data: GlossaryItemModel):
    update_glossary_term(
        term_id=term_id,
        source_term=data.source_term,
        target_term=data.target_term,
        target_lang=data.target_lang or "en",
        category=data.category or "通用",
        is_enabled=data.is_enabled if data.is_enabled is not None else 1
    )
    return {"status": "success", "message": "词条已更新"}

@router.delete("/{term_id}")
def remove_term(term_id: int):
    delete_glossary_term(term_id)
    return {"status": "success", "message": "词条已删除"}

@router.get("/categories")
def get_categories():
    terms = get_glossary_terms()
    cats = sorted(list(set(t.get("category", "通用") for t in terms if t.get("category"))))
    return {"status": "success", "data": ["全部", "通用"] + [c for c in cats if c != "通用"]}

@router.post("/import")
async def import_glossary(file: UploadFile = File(...)):
    """从 Excel (.xlsx) 或 CSV 批量导入术语"""
    filename = file.filename.lower()
    contents = await file.read()
    count = 0

    if filename.endswith(".xlsx"):
        wb = openpyxl.load_workbook(io.BytesIO(contents))
        ws = wb.active
        for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if r_idx == 0:
                # 检查是否表头
                if row and row[0] and ("中文" in str(row[0]) or "Source" in str(row[0])):
                    continue
            if not row or not row[0] or not row[1]:
                continue
            src = str(row[0]).strip()
            tgt = str(row[1]).strip()
            lang = str(row[2]).strip() if len(row) > 2 and row[2] else "en"
            cat = str(row[3]).strip() if len(row) > 3 and row[3] else "通用"
            if src and tgt:
                add_glossary_term(src, tgt, lang, cat)
                count += 1
        wb.close()
    elif filename.endswith(".csv"):
        import csv
        text_stream = io.StringIO(contents.decode("utf-8-sig", errors="ignore"))
        reader = csv.reader(text_stream)
        for r_idx, row in enumerate(reader):
            if r_idx == 0 and ("中文" in str(row[0]) or "Source" in str(row[0])):
                continue
            if not row or len(row) < 2:
                continue
            src = row[0].strip()
            tgt = row[1].strip()
            lang = row[2].strip() if len(row) > 2 and row[2] else "en"
            cat = row[3].strip() if len(row) > 3 and row[3] else "通用"
            if src and tgt:
                add_glossary_term(src, tgt, lang, cat)
                count += 1
    else:
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 或 .csv 格式文件导入")

    return {"status": "success", "imported_count": count, "message": f"成功导入 {count} 条术语"}

@router.get("/export")
def export_glossary():
    """导出全部术语为 Excel"""
    terms = get_glossary_terms()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "制造业术语表"

    # 表头
    headers = ["中文源词条 (Source)", "标准外文译文 (Target)", "适用语种 (Lang)", "工艺分类 (Category)"]
    ws.append(headers)

    # 填充数据
    for t in terms:
        ws.append([t["source_term"], t["target_term"], t["target_lang"], t["category"]])

    # 简单美化列宽
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(15, max_len + 4)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    headers_resp = {
        "Content-Disposition": "attachment; filename=Manufacturing_Glossary.xlsx"
    }
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers_resp)
