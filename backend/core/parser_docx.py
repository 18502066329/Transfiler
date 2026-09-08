import os
import copy
from typing import List, Dict, Any, Tuple
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

class DocxParserEngine:
    """Word (.docx) 高保真解析与双语排版重构引擎"""

    @staticmethod
    def extract_content(file_path: str) -> List[Dict[str, Any]]:
        """
        解析 Word 文档，提取所有可翻译段落和表格单元格
        """
        doc = Document(file_path)
        items = []

        # 1. 提取正文独立段落
        p_count = 0
        for p_idx, paragraph in enumerate(doc.paragraphs):
            text = paragraph.text.strip()
            if text:
                items.append({
                    "id": f"p_{p_idx}",
                    "type": "paragraph",
                    "index": p_idx,
                    "source_text": text,
                    "target_text": "",
                    "matched_terms": []
                })
                p_count += 1

        # 2. 提取表格单元格文本 (按底层元素及多段落粒度提取，避免 GC 导致单元格丢失)
        for t_idx, table in enumerate(doc.tables):
            seen_tcs = set()
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    if cell._tc in seen_tcs:
                        continue
                    seen_tcs.add(cell._tc)

                    for p_idx, paragraph in enumerate(cell.paragraphs):
                        text = paragraph.text.strip()
                        if text:
                            items.append({
                                "id": f"t_{t_idx}_r_{r_idx}_c_{c_idx}_p_{p_idx}",
                                "type": "table_cell",
                                "table_idx": t_idx,
                                "row_idx": r_idx,
                                "col_idx": c_idx,
                                "p_idx": p_idx,
                                "source_text": text,
                                "target_text": "",
                                "matched_terms": []
                            })

        return items

    @staticmethod
    def rebuild_bilingual_doc(
        original_file_path: str,
        output_file_path: str,
        items_map: Dict[str, str],
        layout_mode: str = "zh_top",
        source_lang: str = "zh",
        target_lang: str = "en",
        font_en: str = "Calibri"
    ) -> str:
        """
        高保真重构 Word 文档，生成双语对照文档
        layout_mode:
            - 'zh_top' / 'bilingual_stacked': 中文在上，外文在下 (推荐标准)
            - 'zh_bottom': 中文在下，外文在上
            - 'replace': 纯单语替换
        """
        doc = Document(original_file_path)

        # 规范化 layout_mode
        if layout_mode == "bilingual_stacked":
            layout_mode = "zh_top"

        # 1. 处理独立段落
        for p_idx, paragraph in enumerate(doc.paragraphs):
            item_id = f"p_{p_idx}"
            trans_text = items_map.get(item_id, "").strip()
            if not trans_text:
                continue

            orig_text = paragraph.text
            ref_size = Pt(10.5)
            ref_color = None
            if paragraph.runs:
                if paragraph.runs[0].font.size:
                    ref_size = paragraph.runs[0].font.size
                ref_color = paragraph.runs[0].font.color.rgb if paragraph.runs[0].font.color else None

            if layout_mode == "replace":
                paragraph.text = trans_text
            elif layout_mode == "zh_top":
                if source_lang == "zh":
                    # 中文在原文，外文在译文：中文在上，外文追加在下
                    trans_run = paragraph.add_run(f"\n{trans_text}")
                    trans_run.font.name = font_en
                    trans_run.font.size = ref_size
                    if ref_color:
                        trans_run.font.color.rgb = ref_color
                else:
                    # 外文在原文，中文在译文：中文在上，外文在下
                    paragraph.text = trans_text
                    trans_run = paragraph.add_run(f"\n{orig_text}")
                    trans_run.font.name = font_en
                    trans_run.font.size = ref_size
                    if ref_color:
                        trans_run.font.color.rgb = ref_color
            elif layout_mode == "zh_bottom":
                if source_lang == "zh":
                    # 中文在原文，外文在译文：外文在上，中文在下
                    paragraph.text = trans_text
                    trans_run = paragraph.add_run(f"\n{orig_text}")
                    trans_run.font.size = ref_size
                    if ref_color:
                        trans_run.font.color.rgb = ref_color
                else:
                    # 外文在原文，中文在译文：外文在上，中文追加在下
                    trans_run = paragraph.add_run(f"\n{trans_text}")
                    trans_run.font.size = ref_size
                    if ref_color:
                        trans_run.font.color.rgb = ref_color

        # 2. 处理表格内单元格
        for t_idx, table in enumerate(doc.tables):
            seen_tcs = set()
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    if cell._tc in seen_tcs:
                        continue
                    seen_tcs.add(cell._tc)

                    if not cell.paragraphs:
                        item_id = f"t_{t_idx}_r_{r_idx}_c_{c_idx}"
                        trans_text = items_map.get(item_id, "").strip()
                        if trans_text:
                            p = cell.add_paragraph()
                            p.text = trans_text
                        continue

                    for p_idx, p in enumerate(cell.paragraphs):
                        item_id = f"t_{t_idx}_r_{r_idx}_c_{c_idx}_p_{p_idx}"
                        trans_text = items_map.get(item_id, "").strip()
                        # 兼容单单元格整体匹配
                        if not trans_text and p_idx == 0:
                            trans_text = items_map.get(f"t_{t_idx}_r_{r_idx}_c_{c_idx}", "").strip()
                        if not trans_text:
                            continue

                        orig_text = p.text
                        ref_size = Pt(9.5)
                        ref_color = None
                        if p.runs and p.runs[0].font.size:
                            ref_size = p.runs[0].font.size
                        if p.runs and p.runs[0].font.color:
                            ref_color = p.runs[0].font.color.rgb

                        if layout_mode == "replace":
                            p.text = trans_text
                        elif layout_mode == "zh_top":
                            if source_lang == "zh":
                                # 中文在上，外文在下
                                trans_run = p.add_run(f"\n{trans_text}")
                                trans_run.font.name = font_en
                                trans_run.font.size = ref_size
                                if ref_color:
                                    trans_run.font.color.rgb = ref_color
                            else:
                                # 外文在原文，中文在译文：中文在上，外文在下
                                p.text = trans_text
                                trans_run = p.add_run(f"\n{orig_text}")
                                trans_run.font.name = font_en
                                trans_run.font.size = ref_size
                                if ref_color:
                                    trans_run.font.color.rgb = ref_color
                        elif layout_mode == "zh_bottom":
                            if source_lang == "zh":
                                # 中文在原文，外文在译文：外文在上，中文在下
                                p.text = trans_text
                                trans_run = p.add_run(f"\n{orig_text}")
                                trans_run.font.size = ref_size
                                if ref_color:
                                    trans_run.font.color.rgb = ref_color
                            else:
                                # 外文在原文，中文在译文：外文在上，中文追加在下
                                trans_run = p.add_run(f"\n{trans_text}")
                                trans_run.font.size = ref_size
                                if ref_color:
                                    trans_run.font.color.rgb = ref_color

        # 保存重构后的双语文档
        out_dir = os.path.dirname(output_file_path)
        if out_dir:
            os.makedirs(out_dir, exist_ok=True)
        doc.save(output_file_path)
        return output_file_path
