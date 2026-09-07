import docx
import openpyxl

def inspect():
    # 1. Word 文档检查
    doc = docx.Document('exported_docs/SOP_注塑机标准作业指导书_双语对照版_EN.docx')
    print('=== Word Document Inspection ===')
    for i, p in enumerate(doc.paragraphs[:4]):
        print(f'P{i}: {repr(p.text)}')

    t = doc.tables[0]
    print(f'Table cell (1,0): {repr(t.cell(1,0).text)}')
    print(f'Table cell (1,1): {repr(t.cell(1,1).text)}')

    # 2. Excel 文档检查
    print('\n=== Excel Document Inspection ===')
    wb = openpyxl.load_workbook('exported_docs/工艺参数表_模具点检_双语对照版_EN.xlsx')
    ws = wb.active
    print(f'A3 Value: {repr(ws["A3"].value)}')
    print(f'A3 WrapText: {ws["A3"].alignment.wrap_text}')
    print(f'Row 3 Height: {ws.row_dimensions[3].height}')
    print(f'B3 Value: {repr(ws["B3"].value)}')
    print(f'Row 4 Height: {ws.row_dimensions[4].height}')
    wb.close()

if __name__ == '__main__':
    inspect()
