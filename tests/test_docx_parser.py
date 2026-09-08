import os
import pytest
from pathlib import Path
from backend.core.parser_docx import DocxParserEngine
import docx

def test_docx_extract_and_rebuild(tmp_path):
    sample_path = Path("sample_docs/SOP_注塑机标准作业指导书.docx")
    assert sample_path.exists()

    # 1. 提取内容
    items = DocxParserEngine.extract_content(str(sample_path))
    assert len(items) > 0

    # 验证提取到的条目中包含标题和表格
    has_paragraph = any(it["type"] == "paragraph" for it in items)
    has_table = any(it["type"] == "table_cell" for it in items)
    assert has_paragraph
    assert has_table

    # 2. 模拟翻译结果
    items_map = {}
    for it in items:
        items_map[it["id"]] = f"Translated: {it['source_text']}"

    # 3. 测试中文在上 (zh_top) 重构 Word
    output_docx_top = tmp_path / "test_bilingual_sop_top.docx"
    DocxParserEngine.rebuild_bilingual_doc(
        original_file_path=str(sample_path),
        output_file_path=str(output_docx_top),
        items_map=items_map,
        layout_mode="zh_top",
        source_lang="zh",
        target_lang="en"
    )
    assert output_docx_top.exists()
    doc_top = docx.Document(str(output_docx_top))
    all_text_top = "\n".join([p.text for p in doc_top.paragraphs])
    assert "注塑车间标准作业指导书" in all_text_top
    assert "Translated:" in all_text_top

    # 4. 测试中文在下 (zh_bottom) 重构 Word
    output_docx_bottom = tmp_path / "test_bilingual_sop_bottom.docx"
    DocxParserEngine.rebuild_bilingual_doc(
        original_file_path=str(sample_path),
        output_file_path=str(output_docx_bottom),
        items_map=items_map,
        layout_mode="zh_bottom",
        source_lang="zh",
        target_lang="en"
    )
    assert output_docx_bottom.exists()
    doc_bottom = docx.Document(str(output_docx_bottom))
    first_p = doc_bottom.paragraphs[0].text
    # 中文在下：译文在第一行，原中文在第二行
    assert "Translated:" in first_p

def test_docx_complex_table_extraction_and_rebuild(tmp_path):
    sample_path = Path("sample_docs/职位说明书（体系专员）.docx")
    if not sample_path.exists():
        pytest.skip(f"Sample file {sample_path} not found")

    # 1. 提取内容：验证复杂表格与多段落无损全量提取（原先因 GC 碰撞仅提取 43 条，现应达到 162 条）
    items = DocxParserEngine.extract_content(str(sample_path))
    assert len(items) >= 160, f"Expected at least 160 items extracted, got {len(items)}"

    # 2. 模拟翻译结果字典
    items_map = {it["id"]: f"EN: {it['source_text']}" for it in items}

    # 3. 测试双语重构
    output_docx = tmp_path / "test_bilingual_zhiwei.docx"
    DocxParserEngine.rebuild_bilingual_doc(
        original_file_path=str(sample_path),
        output_file_path=str(output_docx),
        items_map=items_map,
        layout_mode="zh_top",
        source_lang="zh",
        target_lang="en"
    )
    assert output_docx.exists()

    # 4. 验证重构后文档结构完整
    doc = docx.Document(str(output_docx))
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 60

