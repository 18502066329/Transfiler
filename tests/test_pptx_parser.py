import os
import pytest
from pathlib import Path
import pptx
from pptx import Presentation
from backend.core.parser_pptx import PptxParserEngine

def test_pptx_extract_and_rebuild(tmp_path):
    sample_path = Path("sample_docs/注塑车间生产工艺与安全培训课件.pptx")
    assert sample_path.exists(), "测试样本 PPTX 文件必须存在"

    # 1. 提取内容
    items = PptxParserEngine.extract_content(str(sample_path))
    assert len(items) > 0, "提取到的条目数必须大于 0"

    # 验证提取到标题、表格、段落与备注
    has_shape_text = any(it["shape_type"] == "shape_text" for it in items)
    has_table_cell = any(it["shape_type"] == "table_cell" for it in items)
    has_notes = any(it["shape_type"] == "notes_text" for it in items)
    assert has_shape_text, "必须成功提取到文本框形状文字"
    assert has_table_cell, "必须成功提取到表格单元格文字"
    assert has_notes, "必须成功提取到备注页文字"

    # 2. 模拟翻译结果字典
    items_map = {it["id"]: f"EN: {it['source_text']}" for it in items}

    # 3. 测试中文在上 (zh_top) 重构 PPTX 与效果核验
    output_pptx_top = tmp_path / "test_bilingual_top.pptx"
    PptxParserEngine.rebuild_bilingual_doc(
        original_file_path=str(sample_path),
        output_file_path=str(output_pptx_top),
        items_map=items_map,
        layout_mode="zh_top",
        source_lang="zh",
        target_lang="en"
    )
    assert output_pptx_top.exists()

    # 自动化质量核验
    verify_res = PptxParserEngine.verify_pptx_quality(str(output_pptx_top))
    assert verify_res["status"] == "verified"
    assert verify_res["total_slides"] == 3

    # 验证双语内容
    prs_top = Presentation(str(output_pptx_top))
    slide1_text = ""
    for s in prs_top.slides[0].shapes:
        if s.has_text_frame:
            slide1_text += s.text_frame.text + "\n"
    assert "注塑车间生产工艺与安全规程培训" in slide1_text
    assert "EN:" in slide1_text

    # 4. 测试中文在下 (zh_bottom) 重构 PPTX
    output_pptx_bottom = tmp_path / "test_bilingual_bottom.pptx"
    PptxParserEngine.rebuild_bilingual_doc(
        original_file_path=str(sample_path),
        output_file_path=str(output_pptx_bottom),
        items_map=items_map,
        layout_mode="zh_bottom",
        source_lang="zh",
        target_lang="en"
    )
    assert output_pptx_bottom.exists()
    prs_bottom = Presentation(str(output_pptx_bottom))
    assert len(prs_bottom.slides) == 3

    # 5. 测试纯替换 (replace) 重构 PPTX
    output_pptx_replace = tmp_path / "test_bilingual_replace.pptx"
    PptxParserEngine.rebuild_bilingual_doc(
        original_file_path=str(sample_path),
        output_file_path=str(output_pptx_replace),
        items_map=items_map,
        layout_mode="replace",
        source_lang="zh",
        target_lang="en"
    )
    assert output_pptx_replace.exists()
